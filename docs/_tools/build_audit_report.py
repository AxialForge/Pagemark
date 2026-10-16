#!/usr/bin/env python3
"""Build the audit report as a PORTRAIT PDF and a Word (.docx) file from docs/release-audit/*.md.

Standard final step of every /prerelease-audit run (and of any audit/readiness report): the Markdown files are the source of truth,
this turns them into the two documents the owner reads and forwards.

Usage:
  python build_audit_report.py <project-root> --title "Pagemark 1.0.0 gate audit" --version 0.13.0 --commit b1bbf51 \
         [--findings-from "1.0.0 gate audit"] [--out-name Pagemark-1.0.0-Gate-Audit-Report] [--out-dir Documentation] [--rev 1.0]

Needs: python-docx (docx), markdown, and Microsoft Edge or Chrome for the PDF. Never landscape (house rule).
Reads, in order: RELEASE_READINESS.md, FINDINGS.md (from the heading given by --findings-from, else the whole file),
AUDIT_STATUS.md, MANUAL_TEST_PLAN.md. Missing files are skipped.
"""
import argparse, datetime, html, pathlib, re, subprocess, sys, tempfile

try:
    import markdown
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm
except ImportError as e:
    sys.exit(f"missing module: {e}. pip install python-docx markdown")

EDGES = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
         r"C:\Program Files\Google\Chrome\Application\chrome.exe", "/usr/bin/chromium", "/usr/bin/google-chrome"]
SEV = {"BLOCKER": "C00000", "CRITICAL": "E06666", "HIGH": "F4A460", "MEDIUM": "FFD966", "LOW": "D9EAD3", "INFO": "DDEBF7"}
FINDING_COLS = ["ID", "Sev", "Component", "Finding", "Evidence", "Status"]


def browser():
    for p in EDGES:
        if pathlib.Path(p).exists():
            return p
    sys.exit("no Edge/Chrome found for the PDF")


def read(path):
    return path.read_text(encoding="utf-8") if path.exists() else ""


def from_heading(text, heading):
    if not heading:
        return text
    m = re.search(r"^##\s+.*" + re.escape(heading) + r".*$", text, re.M | re.I)
    return text[m.start():] if m else text


def inline(par, text):
    """**bold**, `code`, plain."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        else:
            par.add_run(tok)


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def table_to_docx(doc, rows):
    header = rows[0]
    keep = [header.index(c) for c in FINDING_COLS if c in header] if "Sev" in header and "ID" in header else list(range(len(header)))
    t = doc.add_table(rows=1, cols=len(keep)); t.style = "Table Grid"
    widths = {"ID": 1.7, "Sev": 1.6, "Component": 2.6, "Finding": 6.6, "Evidence": 2.6, "Status": 1.9}
    for i, k in enumerate(keep):
        c = t.rows[0].cells[i]; c.text = ""; inline(c.paragraphs[0], header[k]); shade(c, "D9D9D9")
        for r in c.paragraphs[0].runs: r.bold = True; r.font.size = Pt(8)
    for row in rows[1:]:
        cells = t.add_row().cells
        for i, k in enumerate(keep):
            txt = row[k] if k < len(row) else ""
            cells[i].text = ""; inline(cells[i].paragraphs[0], txt)
            for r in cells[i].paragraphs[0].runs: r.font.size = Pt(8)
            if header[k] == "Sev" and txt.upper() in SEV: shade(cells[i], SEV[txt.upper()])
    if "Sev" in header:
        for row in t.rows:
            for i, k in enumerate(keep):
                row.cells[i].width = Cm(widths.get(header[k], 2.5))
    doc.add_paragraph()


def md_to_docx(doc, text, base=None):
    """base: folder that relative image paths in <img src="..."> lines are resolved against (None = ignore pictures)."""
    lines = text.splitlines(); i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.lstrip().startswith("<"):  # raw HTML: pictures in a "shots" row become real pictures, anything else is skipped
            imgs = re.findall(r'<img src="([^"]+)" alt="([^"]*)"', ln)
            if imgs and base is not None:
                par = doc.add_paragraph(); caps = []
                for src, alt in imgs:
                    path = (pathlib.Path(base) / src).resolve()
                    if path.exists():
                        par.add_run().add_picture(str(path), width=Cm(5.2 if len(imgs) > 1 else 6.5)); par.add_run("  "); caps.append(alt)
                if caps:
                    cp = doc.add_paragraph(" · ".join(caps)); cp.runs[0].italic = True; cp.runs[0].font.size = Pt(8)
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1]):
            rows = [split_row(ln)]; i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            table_to_docx(doc, rows); continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)), level=min(len(m.group(1)), 3))
        elif re.match(r"^\s*[-*]\s+\[[ x]\]\s", ln):
            p = doc.add_paragraph(style="List Bullet"); inline(p, ("☑ " if "[x]" in ln else "☐ ") + re.sub(r"^\s*[-*]\s+\[[ x]\]\s", "", ln))
        elif re.match(r"^\s*[-*]\s+", ln):
            inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", ln))
        elif ln.strip() in ("---", "***"):
            pass
        elif ln.strip():
            inline(doc.add_paragraph(), ln.strip())
        i += 1


def build_docx(out, title, meta, sections, base=None):
    doc = Document()
    sec = doc.sections[0]; sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    for _ in range(6): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    for line in meta:
        q = doc.add_paragraph(); rr = q.add_run(line); rr.font.size = Pt(12)
    doc.add_page_break()
    for name, text in sections:
        md_to_docx(doc, text, base); doc.add_page_break()
    doc.save(out)


CSS = """@page{size:A4 portrait;margin:16mm}body{font-family:Segoe UI,Calibri,Arial,sans-serif;font-size:10pt;line-height:1.4;color:#222}
h1{font-size:20pt;color:#1f3a5f;border-bottom:2px solid #1f3a5f;padding-bottom:4px}h2{font-size:14pt;color:#1f3a5f;margin-top:18px}h3{font-size:11.5pt}
table{border-collapse:collapse;width:100%;margin:8px 0;font-size:8pt;table-layout:auto}th,td{border:1px solid #bbb;padding:3px 5px;vertical-align:top;word-break:break-word}
th{background:#e8e8e8}code{font-family:Consolas,monospace;font-size:8.5pt;background:#f2f2f2;padding:0 2px}.cover{height:240mm;display:flex;flex-direction:column;justify-content:center}
.cover h1{font-size:30pt;border:none}.pb{page-break-after:always}tr{page-break-inside:avoid}"""


def trim_findings(text):
    """Same columns as the Word report: drop Category and Fix commit from findings tables so 8 columns fit a portrait page."""
    out, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        if lines[i].startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(split_row(lines[i])); i += 1
            hdr = block[0]
            if "Sev" in hdr and "ID" in hdr:
                keep = [hdr.index(c) for c in FINDING_COLS if c in hdr]
                rows = [block[0]] + block[2:]
                out.append("| " + " | ".join(hdr[k] for k in keep) + " |"); out.append("|" + "---|" * len(keep))
                out += ["| " + " | ".join((r[k] if k < len(r) else "") for k in keep) + " |" for r in rows[1:]]
            else:
                out += ["| " + " | ".join(r) + " |" if n != 1 else "|" + "---|" * len(hdr) for n, r in enumerate(block)]
            continue
        out.append(lines[i]); i += 1
    return "\n".join(out)


def build_pdf(out, title, meta, sections):
    cover = f'<div class="cover pb"><h1>{html.escape(title)}</h1>' + "".join(f"<p>{html.escape(m)}</p>" for m in meta) + "</div>"
    body = "".join(markdown.markdown(trim_findings(re.sub(r"\[ \]", "☐", re.sub(r"\[x\]", "☑", t))), extensions=["tables"]) + '<div class="pb"></div>' for _, t in sections)
    page = f"<!doctype html><meta charset=utf-8><title>{html.escape(title)}</title><style>{CSS}</style>{cover}{body}"
    # scratch folder beside the output: the 8.3 short Windows temp path makes the browser print nothing
    with tempfile.TemporaryDirectory(dir=str(pathlib.Path(out).parent), ignore_cleanup_errors=True) as d:
        h = pathlib.Path(d) / "r.html"; h.write_text(page, encoding="utf-8")
        subprocess.run([browser(), "--headless", "--disable-gpu", "--no-pdf-header-footer", f"--user-data-dir={d}/profile", f"--print-to-pdf={out}", h.as_uri()],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        # the browser can exit before it has finished writing, and the scratch folder must outlive it:
        # wait until the file exists and stops growing
        import time
        last = -1
        for _ in range(60):
            size = pathlib.Path(out).stat().st_size if pathlib.Path(out).exists() else -1
            if size > 0 and size == last:
                break
            last = size
            time.sleep(0.5)
        else:
            sys.exit(f"the browser did not write {out}")


def next_rev(out_dir, name):
    """First build is Rev 1.0; each rebuild of the same report is the next minor revision (1.1, 1.2...)."""
    best = None
    for p in out_dir.glob(f"{name} Rev *.*"):
        m = re.search(r" Rev (\d+)\.(\d+)\.", p.name)
        if m:
            v = (int(m.group(1)), int(m.group(2)))
            best = v if best is None or v > best else best
    return "1.0" if best is None else f"{best[0]}.{best[1] + 1}"


def main():
    a = argparse.ArgumentParser()
    a.add_argument("root"); a.add_argument("--title", required=True); a.add_argument("--version", default=""); a.add_argument("--commit", default="")
    a.add_argument("--findings-from", default=""); a.add_argument("--out-name", default=""); a.add_argument("--out-dir", default="Documentation"); a.add_argument("--rev", default=""); a.add_argument("--only", nargs="+", default=[])
    a = a.parse_args()
    root = pathlib.Path(a.root).resolve(); src = root / "docs" / "release-audit"
    sections = [("Readiness", read(src / "RELEASE_READINESS.md")),
                ("Findings", from_heading(read(src / "FINDINGS.md"), a.findings_from)),
                ("Status", read(src / "AUDIT_STATUS.md")),
                ("Manual tests", read(src / "MANUAL_TEST_PLAN.md"))]
    if a.only:
        sections = [(f, read(src / f)) for f in a.only]
    sections = [(n, t) for n, t in sections if t.strip()]
    if not sections:
        sys.exit(f"nothing to report: no files in {src}")
    meta = [x for x in (f"Version {a.version}" if a.version else "", f"Commit {a.commit}" if a.commit else "", datetime.date.today().isoformat()) if x]
    name = a.out_name or re.sub(r"[^A-Za-z0-9]+", "-", a.title).strip("-")
    out_dir = root / a.out_dir; out_dir.mkdir(parents=True, exist_ok=True)
    rev = a.rev or next_rev(out_dir, name)
    stem = f"{name} Rev {rev}"
    build_docx(out_dir / f"{stem}.docx", a.title, meta + [f"Revision {rev}"], sections)
    build_pdf(out_dir / f"{stem}.pdf", a.title, meta + [f"Revision {rev}"], sections)
    print(out_dir / f"{stem}.docx"); print(out_dir / f"{stem}.pdf")


main()
